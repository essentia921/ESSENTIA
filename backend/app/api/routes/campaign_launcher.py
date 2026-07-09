from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Header
from pydantic import BaseModel, validator
from typing import Optional, List
from datetime import datetime
import io
import csv
from openpyxl import load_workbook

from backend.app.core.security import decode_access_token
from db.accounts_db import (
    get_all_user_profiles,
    get_profile_info,
    save_launched_campaign,
    get_user_launched_campaigns,
    get_client_tokens,
    verify_account_ownership
)
from backend.app.services.campaign_launcher import launch_campaign, CampaignLauncherError

router = APIRouter()


class TargetItem(BaseModel):
    keyword: Optional[str] = None
    asin: Optional[str] = None
    match_type: Optional[str] = "EXACT"
    bid: Optional[float] = None


class LaunchRequest(BaseModel):
    account_id: int
    profile_id: str
    campaign_type: str
    campaign_name: str
    ad_group_name: str
    asin: str
    daily_budget: float
    default_bid: float
    targets: Optional[List[TargetItem]] = None
    start_date: Optional[str] = None
    bidding_strategy: Optional[str] = "DOWN_ONLY"
    
    @validator('campaign_type')
    def validate_campaign_type(cls, v):
        allowed = ['keywords', 'product', 'product_exact', 'product_expanded', 'automatica']
        if v.lower() not in allowed:
            raise ValueError(f'campaign_type must be one of: {allowed}')
        return v.lower()
    
    @validator('asin')
    def validate_asin(cls, v):
        if not v or len(v) < 10:
            raise ValueError('Invalid ASIN format')
        return v.upper()
    
    @validator('daily_budget', 'default_bid')
    def validate_positive(cls, v):
        if v <= 0:
            raise ValueError('Value must be positive')
        return v


class LaunchResponse(BaseModel):
    success: bool
    campaign_id: Optional[str] = None
    ad_group_id: Optional[str] = None
    message: Optional[str] = None
    error: Optional[str] = None
    amazon_response: Optional[dict] = None


class ParsedTarget(BaseModel):
    keyword: Optional[str] = None
    asin: Optional[str] = None
    match_type: Optional[str] = None
    bid: Optional[float] = None
    row: int
    valid: bool
    error: Optional[str] = None


class ParseResponse(BaseModel):
    success: bool
    targets: List[ParsedTarget]
    valid_count: int
    error_count: int
    message: Optional[str] = None


def get_current_user(authorization: str = None):
    if not authorization:
        raise HTTPException(status_code=401, detail="Not authenticated")
    
    try:
        if authorization.startswith("Bearer "):
            token = authorization[7:]
        else:
            token = authorization
        
        payload = decode_access_token(token)
        if not payload:
            raise HTTPException(status_code=401, detail="Invalid token")
        
        return {
            "id": int(payload.get("sub")),
            "email": payload.get("email")
        }
    except Exception:
        raise HTTPException(status_code=401, detail="Invalid token")


@router.get("/profiles")
async def get_profiles_for_launcher(authorization: str = Header(None)):
    user = get_current_user(authorization)
    user_id = user["id"]
    
    profiles = get_all_user_profiles(user_id)
    
    result = []
    for p in profiles:
        result.append({
            "id": p.get("id"),
            "account_id": p.get("account_id") or p.get("client_account_id"),
            "account_name": p.get("account_name", ""),
            "profile_id": p.get("profile_id"),
            "country_code": p.get("country_code"),
            "marketplace": p.get("marketplace_string", ""),
            "region": p.get("region"),
            "api_url": p.get("api_url"),
        })
    
    return {"profiles": result}


@router.post("/parse-targets")
async def parse_targets_file(
    file: UploadFile = File(...),
    target_type: str = "keywords",
    authorization: str = Header(None)
):
    user = get_current_user(authorization)
    
    content = await file.read()
    filename = file.filename or ""
    
    rows_data = []
    header = None
    
    if filename.lower().endswith('.xlsx') or filename.lower().endswith('.xls'):
        try:
            wb = load_workbook(filename=io.BytesIO(content), read_only=True)
            ws = wb.active
            for row_num, row in enumerate(ws.iter_rows(values_only=True), 1):
                if row_num == 1:
                    header = [str(col).lower().strip() if col else "" for col in row]
                else:
                    row_values = [str(cell).strip() if cell is not None else "" for cell in row]
                    if any(v for v in row_values):
                        rows_data.append((row_num, row_values))
            wb.close()
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Errore lettura file Excel: {str(e)}")
    else:
        try:
            text = content.decode('utf-8')
        except UnicodeDecodeError:
            try:
                text = content.decode('latin-1')
            except:
                raise HTTPException(status_code=400, detail="Could not decode file")
        
        lines = text.strip().split('\n')
        reader = csv.reader(lines)
        
        for row_num, row in enumerate(reader, 1):
            if row_num == 1:
                header = [col.lower().strip() for col in row]
            else:
                if row and any(cell.strip() for cell in row):
                    rows_data.append((row_num, [cell.strip() for cell in row]))
    
    if not header:
        raise HTTPException(status_code=400, detail="File vuoto o senza intestazioni")
    
    targets = []
    valid_count = 0
    error_count = 0
    
    for row_num, row in rows_data:
        target = ParsedTarget(row=row_num, valid=True)
        
        try:
            row_dict = {header[i]: row[i] if i < len(row) else "" for i in range(len(header))}
            
            if target_type == "keywords":
                keyword = row_dict.get('keyword') or row_dict.get('keywords') or row_dict.get('keywordtext', '')
                if not keyword:
                    raise ValueError("Missing keyword")
                
                target.keyword = keyword
                target.match_type = (row_dict.get('match_type') or row_dict.get('matchtype') or 'EXACT').upper()
                
                if target.match_type not in ['EXACT', 'PHRASE', 'BROAD']:
                    target.match_type = 'EXACT'
                
                bid_str = row_dict.get('bid') or row_dict.get('default_bid') or ''
                if bid_str:
                    target.bid = float(str(bid_str).replace(',', '.'))
                
            else:
                asin = row_dict.get('asin') or row_dict.get('target_asin') or ''
                if not asin:
                    raise ValueError("Missing ASIN")
                
                target.asin = asin.upper()
                
                bid_str = row_dict.get('bid') or row_dict.get('default_bid') or ''
                if bid_str:
                    target.bid = float(str(bid_str).replace(',', '.'))
            
            valid_count += 1
            
        except Exception as e:
            target.valid = False
            target.error = str(e)
            error_count += 1
        
        targets.append(target)
    
    return ParseResponse(
        success=error_count == 0,
        targets=targets,
        valid_count=valid_count,
        error_count=error_count,
        message=f"Parsed {valid_count} valid targets, {error_count} errors" if targets else "No targets found"
    )


@router.post("/launch", response_model=LaunchResponse)
async def launch_campaign_endpoint(
    request: LaunchRequest,
    authorization: str = Header(None)
):
    user = get_current_user(authorization)
    user_id = user["id"]
    
    if not verify_account_ownership(user_id, request.account_id):
        raise HTTPException(status_code=403, detail="Non hai accesso a questo account")
    
    if request.campaign_type in ['keywords', 'product_exact', 'product_expanded']:
        has_valid_targets = request.targets and any(t.keyword or t.asin for t in request.targets)
        if not has_valid_targets:
            raise HTTPException(
                status_code=400, 
                detail="Le campagne manuali richiedono almeno un target (keyword o ASIN)"
            )
    
    profile_info = get_profile_info(request.account_id, request.profile_id)
    if not profile_info:
        raise HTTPException(status_code=404, detail="Profile not found")
    
    tokens = get_client_tokens(request.account_id)
    if not tokens:
        raise HTTPException(status_code=400, detail="Account not connected to Amazon")
    
    api_url = profile_info.get("api_url")
    if not api_url:
        from backend.app.services.amazon import get_api_url_for_country
        api_url = get_api_url_for_country(profile_info.get("country_code", "US"))
    
    marketplace = profile_info.get("country_code", "")
    
    targets_list = None
    if request.targets:
        targets_list = [t.dict() for t in request.targets if t.keyword or t.asin]
    
    request_payload = {
        "account_id": request.account_id,
        "profile_id": request.profile_id,
        "campaign_type": request.campaign_type,
        "campaign_name": request.campaign_name,
        "ad_group_name": request.ad_group_name,
        "asin": request.asin,
        "daily_budget": request.daily_budget,
        "default_bid": request.default_bid,
        "targets": targets_list,
        "start_date": request.start_date,
    }
    
    try:
        result = launch_campaign(
            account_id=request.account_id,
            profile_id=request.profile_id,
            api_url=api_url,
            campaign_type=request.campaign_type,
            campaign_name=request.campaign_name,
            ad_group_name=request.ad_group_name,
            asin=request.asin,
            daily_budget=request.daily_budget,
            default_bid=request.default_bid,
            targets=targets_list,
            start_date=request.start_date,
            bidding_strategy=request.bidding_strategy or "DOWN_ONLY"
        )
        
        save_launched_campaign(
            user_id=user_id,
            account_id=request.account_id,
            profile_id=request.profile_id,
            marketplace=marketplace,
            campaign_type=request.campaign_type,
            campaign_id=result.get("campaign_id"),
            ad_group_id=result.get("ad_group_id"),
            asin=request.asin,
            campaign_name=request.campaign_name,
            daily_budget=request.daily_budget,
            default_bid=request.default_bid,
            status="success",
            request_payload=request_payload,
            response_payload=result,
            error_message=None
        )
        
        return LaunchResponse(
            success=True,
            campaign_id=result.get("campaign_id"),
            ad_group_id=result.get("ad_group_id"),
            message="Campaign created successfully"
        )
        
    except CampaignLauncherError as e:
        save_launched_campaign(
            user_id=user_id,
            account_id=request.account_id,
            profile_id=request.profile_id,
            marketplace=marketplace,
            campaign_type=request.campaign_type,
            campaign_id=None,
            ad_group_id=None,
            asin=request.asin,
            campaign_name=request.campaign_name,
            daily_budget=request.daily_budget,
            default_bid=request.default_bid,
            status="failed",
            request_payload=request_payload,
            response_payload=e.amazon_response,
            error_message=str(e)
        )
        
        return LaunchResponse(
            success=False,
            error=str(e),
            amazon_response=e.amazon_response
        )
    
    except Exception as e:
        save_launched_campaign(
            user_id=user_id,
            account_id=request.account_id,
            profile_id=request.profile_id,
            marketplace=marketplace,
            campaign_type=request.campaign_type,
            campaign_id=None,
            ad_group_id=None,
            asin=request.asin,
            campaign_name=request.campaign_name,
            daily_budget=request.daily_budget,
            default_bid=request.default_bid,
            status="failed",
            request_payload=request_payload,
            response_payload=None,
            error_message=str(e)
        )
        
        return LaunchResponse(
            success=False,
            error=f"Unexpected error: {str(e)}"
        )


@router.get("/history")
async def get_launch_history(authorization: str = Header(None), limit: int = 50):
    user = get_current_user(authorization)
    
    campaigns = get_user_launched_campaigns(user["id"], limit)
    
    for c in campaigns:
        if c.get("launched_at"):
            c["launched_at"] = c["launched_at"].isoformat() if hasattr(c["launched_at"], 'isoformat') else str(c["launched_at"])
        if c.get("created_at"):
            c["created_at"] = c["created_at"].isoformat() if hasattr(c["created_at"], 'isoformat') else str(c["created_at"])
        if c.get("daily_budget"):
            c["daily_budget"] = float(c["daily_budget"])
        if c.get("default_bid"):
            c["default_bid"] = float(c["default_bid"])
    
    return {"campaigns": campaigns}
